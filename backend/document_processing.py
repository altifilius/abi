"""Content validation and strictly local document extraction/OCR."""

from __future__ import annotations

import io
import zipfile
from pathlib import PurePosixPath

from defusedxml import ElementTree
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

from .settings import settings

MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".webp": "image/webp",
}
IMAGE_FORMATS = {
    ".png": "PNG",
    ".jpg": "JPEG",
    ".jpeg": "JPEG",
    ".tif": "TIFF",
    ".tiff": "TIFF",
    ".webp": "WEBP",
}
MAX_TEXT_CHARS = 500_000


class DocumentError(ValueError):
    pass


class ProcessingUnavailable(RuntimeError):
    pass


def _image(content: bytes, extension: str) -> Image.Image:
    try:
        with Image.open(io.BytesIO(content)) as image:
            if image.format != IMAGE_FORMATS[extension]:
                raise DocumentError("Image contents do not match the filename extension")
            if image.width * image.height > settings.max_image_pixels:
                raise DocumentError("Image exceeds the pixel limit")
            if getattr(image, "n_frames", 1) > settings.max_document_pages:
                raise DocumentError("Image exceeds the frame limit")
            image.verify()
        image = Image.open(io.BytesIO(content))
        image.load()
        for frame in range(getattr(image, "n_frames", 1)):
            image.seek(frame)
            if image.width * image.height > settings.max_image_pixels:
                image.close()
                raise DocumentError("Image frame exceeds the pixel limit")
            image.load()
        image.seek(0)
        return image
    except (
        UnidentifiedImageError,
        OSError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ) as exc:
        raise DocumentError("Invalid or unsafe image") from exc


def _office(content: bytes, extension: str) -> dict:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > settings.max_zip_entries:
                raise DocumentError("Office archive exceeds the entry limit")
            total = sum(item.file_size for item in entries)
            if total > settings.max_zip_uncompressed_bytes:
                raise DocumentError("Office archive exceeds the expanded size limit")
            names = set()
            for item in entries:
                path = PurePosixPath(item.filename)
                if (
                    item.filename in names
                    or path.is_absolute()
                    or ".." in path.parts
                    or "\\" in item.filename
                ):
                    raise DocumentError("Unsafe or duplicate Office archive entry")
                names.add(item.filename)
                if item.flag_bits & 1 or item.file_size > settings.max_zip_uncompressed_bytes:
                    raise DocumentError("Encrypted or oversized Office archive entry")
                if item.file_size > max(item.compress_size * 200, 1_000_000):
                    raise DocumentError("Office archive exceeds the compression ratio limit")
                if (
                    any(part.lower() in {"embeddings", "activex"} for part in path.parts)
                    or path.name.lower() == "vbaproject.bin"
                ):
                    raise DocumentError("Embedded executable/macro content is not accepted")
                data = archive.read(item)
                if item.filename.endswith((".xml", ".rels")):
                    ElementTree.fromstring(data)
            main = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
            if "[Content_Types].xml" not in names or main not in names:
                raise DocumentError("Office contents do not match the filename extension")
            types = archive.read("[Content_Types].xml")
            expected = (
                b"wordprocessingml.document.main+xml"
                if extension == ".docx"
                else b"spreadsheetml.sheet.main+xml"
            )
            if expected not in types or b"macroEnabled" in types:
                raise DocumentError("Only non-macro DOCX/XLSX files are accepted")
            return {"archive_entries": len(entries), "expanded_bytes": total}
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("Invalid Office archive or unsafe XML") from exc


def validate_document(content: bytes, name: str) -> tuple[str, dict]:
    extension = PurePosixPath(name.lower()).suffix
    if not content:
        raise DocumentError("Empty documents are not accepted")
    if len(content) > settings.max_upload_bytes:
        raise DocumentError("Document exceeds the upload size limit")
    if extension not in MIME_TYPES:
        raise DocumentError("Supported documents: PDF, DOCX, XLSX, PNG, JPEG, TIFF, WEBP")
    metadata = {"extension": extension}
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise DocumentError("File contents are not PDF")
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise DocumentError("Encrypted PDF files are not accepted")
            pages = len(reader.pages)
            if not 0 < pages <= settings.max_document_pages:
                raise DocumentError("PDF exceeds the page limit or has no pages")
            for page in reader.pages:
                page.extract_text()
            metadata["pages"] = pages
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("Invalid PDF structure") from exc
    elif extension in {".docx", ".xlsx"}:
        if not content.startswith(b"PK\x03\x04"):
            raise DocumentError("File contents are not an Office archive")
        metadata.update(_office(content, extension))
    else:
        image = _image(content, extension)
        metadata.update(
            {"width": image.width, "height": image.height, "frames": getattr(image, "n_frames", 1)}
        )
        image.close()
    return MIME_TYPES[extension], metadata


def _ocr(image: Image.Image) -> str:
    import pytesseract

    try:
        config = ""
        if settings.ocr_data_path is not None:
            path = str(settings.ocr_data_path)
            if '"' in path or "\n" in path or "\r" in path:
                raise ProcessingUnavailable("OCR data path contains unsupported characters")
            config = f'--tessdata-dir "{path}"'
        return pytesseract.image_to_string(
            image, lang=settings.ocr_language, config=config, timeout=settings.ocr_timeout_seconds
        )
    except pytesseract.TesseractNotFoundError as exc:
        raise ProcessingUnavailable(
            "Local Tesseract is not installed; configure OCR before processing scans"
        ) from exc
    except pytesseract.TesseractError as exc:
        raise ProcessingUnavailable(
            "Local Tesseract failed; install the configured OCR language data"
        ) from exc
    except RuntimeError as exc:
        raise DocumentError("Local OCR exceeded its processing timeout") from exc


def extract_document(content: bytes, name: str) -> tuple[str, dict]:
    _, metadata = validate_document(content, name)
    extension = metadata["extension"]
    chunks: list[str] = []
    ocr_pages = 0
    if extension == ".pdf":
        import pypdfium2 as pdfium

        reader = PdfReader(io.BytesIO(content), strict=True)
        document = pdfium.PdfDocument(content)
        text_size = 0
        try:
            for index, source_page in enumerate(reader.pages):
                text = (source_page.extract_text() or "").strip()
                if not text:
                    page = document[index]
                    try:
                        width, height = page.get_size()
                        if width * height * 4 > settings.max_image_pixels:
                            raise DocumentError("PDF page exceeds the OCR pixel limit")
                        bitmap = page.render(scale=2)
                        try:
                            image = bitmap.to_pil()
                            try:
                                text = _ocr(image)
                            finally:
                                image.close()
                        finally:
                            bitmap.close()
                    finally:
                        page.close()
                    ocr_pages += 1
                chunks.append(text)
                text_size += len(text)
                if text_size > MAX_TEXT_CHARS:
                    raise DocumentError("Extracted document text exceeds the limit")
        finally:
            document.close()
    elif extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        chunks.extend(paragraph.text for paragraph in document.paragraphs)
        chunks.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            cells = 0
            for sheet in workbook:
                chunks.append(sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > 200_000:
                        raise DocumentError("Spreadsheet exceeds the cell limit")
                    chunks.append(" | ".join(str(value) for value in row if value is not None))
        finally:
            workbook.close()
    else:
        image = _image(content, extension)
        try:
            for frame in range(getattr(image, "n_frames", 1)):
                image.seek(frame)
                if image.width * image.height > settings.max_image_pixels:
                    raise DocumentError("Image frame exceeds the pixel limit")
                chunks.append(_ocr(image))
                ocr_pages += 1
        finally:
            image.close()
    text = "\n".join(chunks).strip()
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentError("Extracted document text exceeds the limit")
    if not text:
        raise DocumentError("No readable text extracted; document has not been verified")
    return text, {
        **metadata,
        "ocr_pages": ocr_pages,
        "extraction": "local",
        "requires_human_review": True,
    }
