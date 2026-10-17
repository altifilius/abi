"""Validated rich documents and local Unicode exports; user input is never markup."""

from __future__ import annotations

import html
import io
import json
import re
from pathlib import Path
from urllib.parse import urlsplit

from .exports import _FONT_LOCK, ExportUnavailable, clean_text
from .settings import settings

BLOCKS = {"paragraph", "heading", "bulletList", "orderedList", "table"}
MARKS = {"bold", "italic", "underline", "strike", "code", "link"}
MAX_DOCUMENT_BYTES = 1_000_000


def text_document(text: str = "") -> dict:
    return {
        "type": "doc",
        "content": [
            {"type": "paragraph", **({"content": [{"type": "text", "text": line}]} if line else {})}
            for line in text.split("\n")
        ],
    }


def safe_link(value: object) -> str:
    if not isinstance(value, str) or len(value) > 2000 or any(ord(c) < 32 for c in value):
        raise ValueError("Invalid link")
    parsed = urlsplit(value)
    if (
        parsed.scheme in {"http", "https"}
        and parsed.hostname
        and not parsed.username
        and not parsed.password
    ):
        return value
    if parsed.scheme == "mailto" and parsed.path and not parsed.netloc:
        return value
    raise ValueError("Only HTTP(S) and mailto links are supported")


def table_grid(node: dict) -> tuple[list[list[dict | None]], list[tuple[int, int, int, int]]]:
    rows = node.get("content", [])
    occupied: dict[tuple[int, int], dict | None] = {}
    spans = []
    width = 0
    for row_index, row in enumerate(rows):
        column = 0
        for cell in row.get("content", []):
            while (row_index, column) in occupied:
                column += 1
            attrs = cell.get("attrs", {})
            colspan, rowspan = attrs.get("colspan", 1), attrs.get("rowspan", 1)
            if row_index + rowspan > len(rows) or column + colspan > 100:
                raise ValueError("Table span extends beyond the table")
            for y in range(row_index, row_index + rowspan):
                for x in range(column, column + colspan):
                    if (y, x) in occupied:
                        raise ValueError("Overlapping table cells")
                    occupied[y, x] = cell if (y, x) == (row_index, column) else None
            if colspan > 1 or rowspan > 1:
                spans.append((column, row_index, column + colspan - 1, row_index + rowspan - 1))
            column += colspan
            width = max(width, column)
    if not width or any((y, x) not in occupied for y in range(len(rows)) for x in range(width)):
        raise ValueError("Table rows must form a rectangular grid")
    return [[occupied[y, x] for x in range(width)] for y in range(len(rows))], spans


def validate_document(document: object) -> dict:
    """Reject unknown nodes/attributes instead of silently dropping user formatting."""
    try:
        size = len(json.dumps(document, ensure_ascii=False, allow_nan=False).encode("utf-8"))
    except (ValueError, TypeError, UnicodeError) as error:
        raise ValueError("Document must be valid Unicode JSON") from error
    if size > MAX_DOCUMENT_BYTES:
        raise ValueError("Document is too large")
    count = 0

    def visit(node: object, parent: str | None = None, depth: int = 0) -> None:
        nonlocal count
        count += 1
        if count > 20_000 or depth > 16 or not isinstance(node, dict):
            raise ValueError("Document structure exceeds supported limits")
        kind = node.get("type")
        if not isinstance(kind, str):
            raise ValueError("Document node type must be text")
        allowed_children = {
            None: {"doc"},
            "doc": BLOCKS,
            "paragraph": {"text", "hardBreak"},
            "heading": {"text", "hardBreak"},
            "bulletList": {"listItem"},
            "orderedList": {"listItem"},
            "listItem": BLOCKS - {"table"},
            "table": {"tableRow"},
            "tableRow": {"tableCell", "tableHeader"},
            "tableCell": BLOCKS - {"table"},
            "tableHeader": BLOCKS - {"table"},
        }
        if kind not in allowed_children.get(parent, set()):
            raise ValueError("Unsupported document node or nesting")
        if set(node) - {"type", "content", "attrs", "text", "marks"}:
            raise ValueError("Unsupported document properties")
        attrs = node.get("attrs", {})
        if not isinstance(attrs, dict):
            raise ValueError("Node attributes must be an object")
        allowed_attrs = {
            "heading": {"level", "textAlign"},
            "paragraph": {"textAlign"},
            "orderedList": {"start"},
            "tableCell": {"colspan", "rowspan", "colwidth"},
            "tableHeader": {"colspan", "rowspan", "colwidth"},
        }.get(kind, set())
        if set(attrs) - allowed_attrs:
            raise ValueError("Unsupported node attributes")
        if kind == "heading" and (
            type(attrs.get("level")) is not int or not 1 <= attrs["level"] <= 6
        ):
            raise ValueError("Heading level must be 1–6")
        if attrs.get("textAlign") not in (None, "left", "center", "right", "justify"):
            raise ValueError("Unsupported paragraph alignment")
        for key, maximum in (("start", 1_000_000), ("colspan", 20), ("rowspan", 20)):
            if key in attrs and (type(attrs[key]) is not int or not 1 <= attrs[key] <= maximum):
                raise ValueError(f"Invalid {key}")
        widths = attrs.get("colwidth")
        if widths is not None and (
            not isinstance(widths, list)
            or len(widths) != attrs.get("colspan", 1)
            or any(type(width) is not int or not 1 <= width <= 5000 for width in widths)
        ):
            raise ValueError("Invalid table column widths")
        if kind == "text":
            text = node.get("text")
            if not isinstance(text, str) or not text or clean_text(text) != text:
                raise ValueError("Text must be nonempty Unicode without control characters")
            if "content" in node or attrs:
                raise ValueError("Text nodes cannot have children or attributes")
        elif "text" in node or "marks" in node:
            raise ValueError("Only text nodes may contain text or marks")
        marks = node.get("marks", [])
        if not isinstance(marks, list) or len(marks) > len(MARKS):
            raise ValueError("Invalid text marks")
        used_marks = set()
        for mark in marks:
            if not isinstance(mark, dict) or set(mark) - {"type", "attrs"}:
                raise ValueError("Invalid text mark")
            name = mark.get("type")
            if not isinstance(name, str):
                raise ValueError("Text mark type must be text")
            if name not in MARKS or name in used_marks:
                raise ValueError("Unsupported or duplicate text mark")
            used_marks.add(name)
            values = mark.get("attrs", {})
            if not isinstance(values, dict):
                raise ValueError("Invalid mark attributes")
            if name == "link":
                if (
                    set(values) - {"href", "target", "rel", "class"}
                    or values.get("class") is not None
                ):
                    raise ValueError("Unsupported link attributes")
                safe_link(values.get("href"))
                if values.get("target") not in (None, "_blank", "_self"):
                    raise ValueError("Unsupported link target")
                if values.get("rel") is not None and values["rel"] not in (
                    "",
                    "noopener noreferrer",
                    "noopener noreferrer nofollow",
                    "nofollow",
                ):
                    raise ValueError("Unsupported link relationship")
            elif values:
                raise ValueError("Unsupported mark attributes")
        children = node.get("content", [])
        if not isinstance(children, list) or (kind in {"text", "hardBreak"} and children):
            raise ValueError("Invalid node children")
        if (
            kind
            in {
                "bulletList",
                "orderedList",
                "listItem",
                "table",
                "tableCell",
                "tableHeader",
            }
            and not children
        ):
            raise ValueError("Lists and tables require content")
        if kind == "table" and len(children) > 100:
            raise ValueError("A table supports at most 100 rows")
        for child in children:
            visit(child, kind, depth + 1)
        if kind == "table":
            table_grid(node)

    visit(document)
    return document


def render_docx(title: str, sections: list[dict], notice: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    document.styles["Normal"].font.name = "DejaVu Sans"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(clean_text(title), 0)
    document.add_paragraph(clean_text(notice))

    def inline(paragraph, node, header=False):
        if node["type"] == "hardBreak":
            paragraph.add_run().add_break()
            return
        marks = {mark["type"]: mark.get("attrs", {}) for mark in node.get("marks", [])}
        run = paragraph.add_run(node["text"])
        run.bold = True if header or "bold" in marks else None
        run.italic = True if "italic" in marks else None
        run.underline = True if "underline" in marks else None
        run.font.strike = True if "strike" in marks else None
        if "code" in marks:
            run.font.name = "DejaVu Sans Mono"
        if "link" in marks:
            relation = paragraph.part.relate_to(
                marks["link"]["href"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relation)
            run.underline = True
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)

    def blocks(container, nodes, depth=0, list_id=None, header=False):
        for node in nodes:
            kind = node["type"]
            if kind in {"paragraph", "heading"}:
                style = f"Heading {node['attrs']['level']}" if kind == "heading" else None
                paragraph = container.add_paragraph(style=style)
                alignment = node.get("attrs", {}).get("textAlign")
                if alignment:
                    paragraph.alignment = {
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT,
                        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    }[alignment]
                if depth:
                    paragraph.paragraph_format.left_indent = Inches(0.25 * depth)
                if list_id is not None:
                    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    numbering.get_or_add_ilvl().val = 0
                    numbering.get_or_add_numId().val = list_id
                    list_id = None
                for child in node.get("content", []):
                    inline(paragraph, child, header)
            elif kind in {"bulletList", "orderedList"}:
                numbering = document.part.numbering_part.element
                abstract_id = (
                    max(
                        [
                            int(item.get(qn("w:abstractNumId")))
                            for item in numbering.findall(qn("w:abstractNum"))
                        ]
                        + [-1]
                    )
                    + 1
                )
                number_id = (
                    max(
                        [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
                        + [0]
                    )
                    + 1
                )
                abstract = OxmlElement("w:abstractNum")
                abstract.set(qn("w:abstractNumId"), str(abstract_id))
                level = OxmlElement("w:lvl")
                level.set(qn("w:ilvl"), "0")
                for tag, value in (
                    ("start", str(node.get("attrs", {}).get("start", 1))),
                    ("numFmt", "decimal" if kind == "orderedList" else "bullet"),
                    ("lvlText", "%1." if kind == "orderedList" else "•"),
                ):
                    element = OxmlElement(f"w:{tag}")
                    element.set(qn("w:val"), value)
                    level.append(element)
                abstract.append(level)
                numbering.insert(0, abstract)
                number = OxmlElement("w:num")
                number.set(qn("w:numId"), str(number_id))
                reference = OxmlElement("w:abstractNumId")
                reference.set(qn("w:val"), str(abstract_id))
                number.append(reference)
                numbering.append(number)
                for item in node["content"]:
                    blocks(container, item["content"], depth + 1, number_id, header)
            elif kind == "table":
                grid, spans = table_grid(node)
                table = container.add_table(rows=len(grid), cols=len(grid[0]))
                table.style = "Table Grid"
                for x, y, end_x, end_y in spans:
                    table.cell(y, x).merge(table.cell(end_y, end_x))
                for y, row in enumerate(grid):
                    for x, cell in enumerate(row):
                        if cell is not None:
                            target = table.cell(y, x)
                            first = target.paragraphs[0]
                            target._tc.remove(first._p)
                            blocks(target, cell["content"], header=cell["type"] == "tableHeader")
                            widths = cell.get("attrs", {}).get("colwidth")
                            if widths:
                                target.width = Inches(sum(widths) / 96)

    for section in sections:
        document.add_heading(clean_text(section["title"]), 1)
        blocks(document, section["content"].get("content", []))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_fonts() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular = Path(settings.pdf_font_path)
    bold = Path(settings.pdf_bold_font_path)
    stem = re.sub(r"(?:-Regular)?$", "", regular.stem)
    italic = regular.with_name(f"{stem}-Oblique.ttf")
    bold_italic = regular.with_name(f"{stem}-BoldOblique.ttf")
    if not italic.is_file():
        italic = regular.with_name(f"{stem}-Italic.ttf")
    if not bold_italic.is_file():
        bold_italic = regular.with_name(f"{stem}-BoldItalic.ttf")
    fonts = {
        "AbiReport": regular,
        "AbiReportBold": bold,
        "AbiReportItalic": italic,
        "AbiReportBoldItalic": bold_italic,
    }
    if any(not path.is_file() for path in fonts.values()):
        raise ExportUnavailable(
            "Configure local Unicode regular, bold, italic and bold-italic PDF fonts"
        )
    with _FONT_LOCK:
        for name, path in fonts.items():
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path)))
        pdfmetrics.registerFontFamily(
            "AbiReport",
            normal="AbiReport",
            bold="AbiReportBold",
            italic="AbiReportItalic",
            boldItalic="AbiReportBoldItalic",
        )


def render_pdf(title: str, sections: list[dict], notice: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _pdf_fonts()
    normal = ParagraphStyle(
        "ReportBody",
        fontName="AbiReport",
        fontSize=10,
        leading=15,
        spaceAfter=5 * mm,
        wordWrap="CJK",
    )
    heading = ParagraphStyle(
        "ReportHeading",
        parent=normal,
        fontName="AbiReportBold",
        fontSize=14,
        leading=20,
        spaceBefore=4 * mm,
    )
    title_style = ParagraphStyle("ReportTitle", parent=heading, fontSize=20, leading=27)

    def escaped(value):
        return html.escape(clean_text(value), quote=True).replace("\n", "<br/>")

    def inline(node):
        if node["type"] == "hardBreak":
            return "<br/>"
        value = escaped(node["text"])
        for mark in node.get("marks", []):
            kind = mark["type"]
            tag = {
                "bold": "b",
                "italic": "i",
                "underline": "u",
                "strike": "strike",
                "code": "font",
            }.get(kind)
            if kind == "link":
                value = f'<link href="{html.escape(mark["attrs"]["href"], quote=True)}" color="#354fc0"><u>{value}</u></link>'
            elif kind == "code":
                value = f'<font backColor="#f1f2f5">{value}</font>'
            else:
                value = f"<{tag}>{value}</{tag}>"
        return value

    def blocks(nodes, available_width=174 * mm, header=False):
        result = []
        for node in nodes:
            kind = node["type"]
            if kind in {"paragraph", "heading"}:
                level = node.get("attrs", {}).get("level", 1)
                base = heading if kind == "heading" else normal
                attrs = node.get("attrs", {})
                style = ParagraphStyle(
                    "ReportNode",
                    parent=base,
                    alignment={
                        "left": TA_LEFT,
                        "center": TA_CENTER,
                        "right": TA_RIGHT,
                        "justify": TA_JUSTIFY,
                    }.get(attrs.get("textAlign"), TA_LEFT),
                )
                if kind == "heading":
                    style.fontSize = max(10, 16 - level)
                if header:
                    style.fontName = "AbiReportBold"
                value = "".join(inline(child) for child in node.get("content", [])) or "&#160;"
                result.append(Paragraph(value, style))
            elif kind in {"bulletList", "orderedList"}:
                items = [
                    ListItem(blocks(item["content"], available_width - 18, header))
                    for item in node["content"]
                ]
                result.append(
                    ListFlowable(
                        items,
                        bulletType="1" if kind == "orderedList" else "bullet",
                        start=node.get("attrs", {}).get("start", 1)
                        if kind == "orderedList"
                        else "circle",
                        leftIndent=18,
                        bulletFontName="AbiReport",
                    )
                )
            elif kind == "table":
                grid, spans = table_grid(node)
                widths = [available_width / len(grid[0])] * len(grid[0])
                explicit = [None] * len(grid[0])
                for row in grid:
                    for x, cell in enumerate(row):
                        if cell and cell.get("attrs", {}).get("colwidth"):
                            for offset, width in enumerate(cell["attrs"]["colwidth"]):
                                explicit[x + offset] = width
                if all(width is not None for width in explicit):
                    total = sum(explicit)
                    widths = [available_width * width / total for width in explicit]
                cells = [
                    [
                        blocks(
                            cell["content"],
                            sum(widths[x : x + cell.get("attrs", {}).get("colspan", 1)]) - 12,
                            cell["type"] == "tableHeader",
                        )
                        if cell
                        else ""
                        for x, cell in enumerate(row)
                    ]
                    for row in grid
                ]
                table = Table(cells, colWidths=widths, hAlign="LEFT", splitByRow=1, splitInRow=1)
                commands = [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b9bdc8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]
                commands.extend(("SPAN", (x, y), (end_x, end_y)) for x, y, end_x, end_y in spans)
                for y, row in enumerate(grid):
                    for x, cell in enumerate(row):
                        if cell and cell["type"] == "tableHeader":
                            commands.append(
                                ("BACKGROUND", (x, y), (x, y), colors.HexColor("#f1f2f5"))
                            )
                table.setStyle(TableStyle(commands))
                result.extend([table, Spacer(1, 4 * mm)])
        return result

    story = [Paragraph(escaped(title), title_style), Paragraph(escaped(notice), normal)]
    for section in sections:
        story.append(Paragraph(escaped(section["title"]), heading))
        story.extend(blocks(section["content"].get("content", [])))
    output = io.BytesIO()
    SimpleDocTemplate(
        output,
        pagesize=(210 * mm, 297 * mm),
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=clean_text(title),
    ).build(story)
    return output.getvalue()


def render_report(format: str, title: str, sections: list[dict], notice: str) -> tuple[bytes, str]:
    for section in sections:
        validate_document(section["content"])
    if format == "docx":
        return render_docx(
            title, sections, notice
        ), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if format == "pdf":
        return render_pdf(title, sections, notice), "application/pdf"
    raise ValueError("Supported export formats: docx, pdf")
