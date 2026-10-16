import hashlib
import io
import json
import tempfile
import unittest
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import HTTPException
from PIL import Image
from pydantic import ValidationError
from pypdf import PdfReader

from backend.applications import (
    ApplicationData,
    _persist,
    _read_document,
    _verified_document_stream,
    approval_hash,
    budget_check,
    require_current_approval,
    require_revision,
    storage_path,
)
from backend.document_processing import DocumentError, extract_document, validate_document
from backend.exports import package_zip, render_docx, render_pdf
from backend.settings import settings


class DocumentValidationTests(unittest.TestCase):
    def png(self):
        output = io.BytesIO()
        Image.new("RGB", (10, 10), "white").save(output, "PNG")
        return output.getvalue()

    def test_image_cannot_masquerade_as_pdf(self):
        with self.assertRaises(DocumentError):
            validate_document(self.png(), "scan.pdf")

    def test_fake_and_misnamed_images_rejected(self):
        for content, name in [(b"not a real image", "photo.png"), (self.png(), "photo.jpg")]:
            with self.subTest(name=name), self.assertRaises(DocumentError):
                validate_document(content, name)

    def test_office_zip_bomb_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"0" * 1_500_000)
            archive.writestr("[Content_Types].xml", "<Types/>")
        with self.assertRaises(DocumentError):
            validate_document(output.getvalue(), "report.docx")

    def test_archive_traversal_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("../escape.xml", "<root/>")
        with self.assertRaises(DocumentError):
            validate_document(output.getvalue(), "report.docx")

    def test_private_content_persists_integrity_checked(self):
        with (
            tempfile.TemporaryDirectory() as root,
            patch.object(settings, "storage_root", Path(root)),
        ):
            content = self.png()
            digest = hashlib.sha256(content).hexdigest()
            key = "workspace/document/" + digest
            _persist(key, content)
            document = SimpleNamespace(storage_key=key, size=len(content), sha256=digest)
            self.assertEqual(_read_document(document), content)
            self.assertEqual(b"".join(_verified_document_stream(document)), content)
            self.assertEqual(storage_path(key).stat().st_mode & 0o777, 0o600)
            storage_path(key).write_bytes(b"altered")
            with self.assertRaises(HTTPException) as raised:
                _read_document(document)
            self.assertEqual(raised.exception.status_code, 409)
            with self.assertRaises(HTTPException):
                storage_path("../escape")

    def test_xlsx_extraction_is_local(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Proje", "İstanbul", 123])
        output = io.BytesIO()
        workbook.save(output)
        text, metadata = extract_document(output.getvalue(), "budget.xlsx")
        self.assertIn("İstanbul", text)
        self.assertEqual(metadata["extraction"], "local")
        self.assertTrue(metadata["requires_human_review"])


class BudgetAndApprovalTests(unittest.TestCase):
    def budget(self, amount="0.10", currency="TRY"):
        return {
            "currency": currency,
            "items": [
                {
                    "id": "p",
                    "category": "personnel",
                    "description": "Research engineer",
                    "amount": amount,
                }
            ],
        }

    def finance(self):
        return {
            "currency": "TRY",
            "eligible_categories": ["personnel"],
            "support_kind": "grant",
            "support_rate": "0.75",
            "max_support": "1000",
        }

    def test_decimal_budget_and_no_invented_exchange(self):
        budget = self.budget()
        budget["items"].append(
            {"id": "p2", "category": "personnel", "description": "Research", "amount": "0.20"}
        )
        result = budget_check(budget, self.finance(), {})
        self.assertEqual(result["budget"]["total"], "0.30")
        mismatch = budget_check(self.budget(currency="EUR"), self.finance(), {})
        self.assertIn("currency_mismatch", [error["code"] for error in mismatch["errors"]])

    def test_non_grant_matching_not_invented(self):
        finance = self.finance() | {"support_kind": "tax_credit", "support_rate": "0.20"}
        result = budget_check(self.budget("100"), finance, {})
        self.assertIsNone(result["budget"]["support_estimate"])
        self.assertIsNone(result["budget"]["own_contribution"])

    def test_employment_incentive_does_not_require_fictional_project_budget(self):
        finance = {
            "currency": "GBP",
            "support_kind": "tax_credit",
            "budget_required": False,
            "eligible_categories": [],
            "max_support": "10500",
        }
        result = budget_check({"currency": "GBP", "items": []}, finance, {})
        self.assertEqual(result["errors"], [])
        self.assertIsNone(result["budget"]["support_estimate"])
        self.assertIsNone(result["budget"]["own_contribution"])

    def test_float_negative_and_unknown_budget_rejected(self):
        for amount in ["-1", "NaN", "1e9", 1.25]:
            with self.subTest(amount=amount), self.assertRaises(ValidationError):
                ApplicationData(budget=self.budget(amount))

    def test_approval_fingerprint_changes_with_sections_and_evaluation(self):
        application = SimpleNamespace(
            data={"sections": {"purpose": "Original"}, "attachments": {}},
            revision=1,
            profile_revision=2,
            call_version=3,
            source_revision_id="source-v3",
            rule_version="v2",
            template_snapshot={"sections": []},
        )
        evaluation = SimpleNamespace(id="evaluation-1")
        document = SimpleNamespace(
            id="document", sha256="a" * 64, verified_by="reviewer", verified_at=datetime.now(UTC)
        )
        original = approval_hash(application, evaluation, [document])
        application.data["sections"]["purpose"] = "Edited actual field"
        self.assertNotEqual(original, approval_hash(application, evaluation, [document]))
        application.data["sections"]["purpose"] = "Original"
        evaluation.id = "new-evaluation"
        self.assertNotEqual(original, approval_hash(application, evaluation, [document]))
        with self.assertRaises(HTTPException) as raised:
            require_revision(application, 99)
        self.assertEqual(raised.exception.status_code, 409)

    def test_stale_approval_cannot_export_changed_sections_or_revision(self):
        application = SimpleNamespace(
            state="approved",
            data={"sections": {"purpose": "Reviewed"}, "attachments": {}},
            revision=1,
            profile_revision=2,
            call_version=3,
            source_revision_id="source",
            rule_version="v1",
            template_snapshot={},
        )
        evaluation = SimpleNamespace(id="evaluated")
        approval = SimpleNamespace(
            approved=True,
            revision=1,
            reviewer_id="human-reviewer",
            content_hash=approval_hash(application, evaluation, []),
        )
        require_current_approval(application, approval, evaluation, [])
        application.data["sections"]["purpose"] = "Changed after approval"
        with self.assertRaises(HTTPException) as raised:
            require_current_approval(application, approval, evaluation, [])
        self.assertEqual(raised.exception.status_code, 409)
        application.data["sections"]["purpose"] = "Reviewed"
        application.revision = 2
        with self.assertRaises(HTTPException):
            require_current_approval(application, approval, evaluation, [])


class RealExportTests(unittest.TestCase):
    def test_manual_docx_roundtrip_and_escaped_unicode(self):
        from docx import Document

        content = render_docx(
            "İstanbul Ar-Ge",
            [("Açıklama", "<script> & Şirket çözümü")],
            notice="General draft, not an official form",
        )
        text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
        self.assertIn("<script> & Şirket çözümü", text)
        self.assertIn("İstanbul", text)
        validate_document(content, "draft.docx")

    def test_real_unicode_pdf_and_package_attachments(self):
        sections = [("Amaç", "İstanbul araştırma <b>markup is literal</b> & Türkçe")]
        pdf = render_pdf("Proje taslağı", sections, notice="Not an official application")
        reader = PdfReader(io.BytesIO(pdf))
        self.assertIn("İstanbul", "".join(page.extract_text() for page in reader.pages))
        package = package_zip(
            "Project",
            sections,
            guide="Submit manually at https://www.tubitak.gov.tr",
            sources=[{"url": "https://www.tubitak.gov.tr"}],
            attachments=[("original.txt", b"actual attachment bytes")],
            manifest={"revision": 1},
        )
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            self.assertEqual(
                archive.read("attachments/001-original.txt"), b"actual attachment bytes"
            )
            self.assertEqual(
                archive.read("SUBMISSION-GUIDE.txt").decode(),
                "Submit manually at https://www.tubitak.gov.tr",
            )
            self.assertEqual(
                json.loads(archive.read("sources.json")), [{"url": "https://www.tubitak.gov.tr"}]
            )
            self.assertEqual(json.loads(archive.read("manifest.json")), {"revision": 1})
            attachment_pdf = PdfReader(io.BytesIO(archive.read("project-preparation.pdf")))
            self.assertIn(
                "İstanbul araştırma", "".join(page.extract_text() for page in attachment_pdf.pages)
            )
            from docx import Document

            attachment_docx = Document(io.BytesIO(archive.read("project-preparation.docx")))
            self.assertIn(
                "İstanbul araştırma <b>markup is literal</b> & Türkçe",
                "\n".join(paragraph.text for paragraph in attachment_docx.paragraphs),
            )


if __name__ == "__main__":
    unittest.main()
