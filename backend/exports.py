"""Local, non-official project drafts and source-linked preparation packages."""

from __future__ import annotations

import html
import io
import json
import re
import threading
import zipfile
from pathlib import Path

from .settings import settings

_FONT_LOCK = threading.Lock()


class ExportUnavailable(RuntimeError):
    pass


def clean_text(value: object) -> str:
    # XML 1.0 disallows control characters; user prose is never interpreted as markup.
    return "".join(
        character for character in str(value) if character in "\n\r\t" or ord(character) >= 32
    ).replace("\x7f", "")


def readable(value: object) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return clean_text(value)


def generic_sections(project) -> list[tuple[str, str]]:
    sections = [("Project description / Proje açıklaması", project.description or "")]
    for group, value in (project.profile or {}).items():
        sections.append((group.replace("_", " ").title(), readable(value)))
    return sections


def render_docx(title: str, sections: list[tuple[str, str]], *, notice: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.styles["Normal"].font.name = "DejaVu Sans"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(clean_text(title), 0)
    document.add_paragraph(clean_text(notice))
    for label, text in sections:
        document.add_heading(clean_text(label), 1)
        for paragraph in clean_text(text).split("\n"):
            document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf(title: str, sections: list[tuple[str, str]], *, notice: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_path = Path(settings.pdf_font_path)
    if not font_path.is_file():
        raise ExportUnavailable("Configure PDF_FONT_PATH to a local Unicode TrueType font")
    with _FONT_LOCK:
        if "AbiUnicode" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("AbiUnicode", str(font_path)))
    normal = ParagraphStyle(
        "AbiBody",
        fontName="AbiUnicode",
        fontSize=10,
        leading=15,
        alignment=TA_LEFT,
        wordWrap="CJK",
        spaceAfter=5 * mm,
    )
    heading = ParagraphStyle(
        "AbiHeading",
        parent=normal,
        fontSize=14,
        leading=19,
        textColor=colors.HexColor("#532888"),
        spaceBefore=4 * mm,
    )
    title_style = ParagraphStyle("AbiTitle", parent=heading, fontSize=20, leading=26)

    def paragraph(text: str, style=normal):
        return Paragraph(html.escape(clean_text(text), quote=False).replace("\n", "<br/>"), style)

    story = [paragraph(title, title_style), paragraph(notice), Spacer(1, 4 * mm)]
    for label, text in sections:
        story.append(paragraph(label, heading))
        # Split long prose into modest blocks so a paragraph cannot overflow a page.
        for line in clean_text(text).split("\n"):
            for start in range(0, max(len(line), 1), 2000):
                story.append(paragraph(line[start : start + 2000]))
    output = io.BytesIO()
    SimpleDocTemplate(
        output,
        pagesize=(210 * mm, 297 * mm),
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=clean_text(title),
    ).build(story)
    return output.getvalue()


def render(
    format: str, title: str, sections: list[tuple[str, str]], notice: str
) -> tuple[bytes, str]:
    if format == "docx":
        return render_docx(
            title, sections, notice=notice
        ), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if format == "pdf":
        return render_pdf(title, sections, notice=notice), "application/pdf"
    raise ValueError("Supported export formats: docx, pdf")


def safe_filename(name: str) -> str:
    name = re.sub(r"[^\w. -]", "_", Path(name).name, flags=re.UNICODE).strip(" .")
    return name[:150] or "document"


def package_zip(
    title: str,
    sections: list[tuple[str, str]],
    *,
    guide: str,
    sources: list[dict],
    attachments: list[tuple[str, bytes]],
    manifest: dict,
) -> bytes:
    if not guide.strip() or not sources:
        raise ExportUnavailable(
            "A source-linked submission guide is required for a preparation package"
        )
    notice = "Source-linked application preparation package. Not an official application form; no portal submission has occurred."
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("project-preparation.docx", render_docx(title, sections, notice=notice))
        archive.writestr("project-preparation.pdf", render_pdf(title, sections, notice=notice))
        archive.writestr("SUBMISSION-GUIDE.txt", clean_text(guide).encode("utf-8"))
        archive.writestr(
            "sources.json", json.dumps(sources, ensure_ascii=False, indent=2).encode("utf-8")
        )
        archive.writestr(
            "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
        )
        for index, (name, content) in enumerate(attachments, 1):
            archive.writestr(f"attachments/{index:03d}-{safe_filename(name)}", content)
    return output.getvalue()
